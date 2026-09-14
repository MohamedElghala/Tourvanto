import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()
# Remove default sheet
wb.remove(wb.active)

# Color Palette: Brand Teal & Luxury Slate
brand_teal = "0D9488"
dark_slate = "0F172A"
header_fill = PatternFill(start_color=dark_slate, end_color=dark_slate, fill_type="solid")
brand_fill = PatternFill(start_color=brand_teal, end_color=brand_teal, fill_type="solid")
accent_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
highlight_fill = PatternFill(start_color="CCFBF1", end_color="CCFBF1", fill_type="solid")

header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
title_font = Font(name="Arial", size=14, bold=True, color="0D9488")
bold_font = Font(name="Arial", size=10, bold=True)
regular_font = Font(name="Arial", size=10)

thin_border_side = Side(border_style="thin", color="CBD5E1")
cell_border = Border(top=thin_border_side, left=thin_border_side, right=thin_border_side, bottom=thin_border_side)

def style_sheet(ws, title_text):
    ws.views.sheetView[0].showGridLines = True
    ws.merge_cells("A1:G1")
    t_cell = ws["A1"]
    t_cell.value = f"Tourvanto International - {title_text}"
    t_cell.font = title_font
    t_cell.alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[1].height = 35

def auto_fit_columns(ws):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val = str(cell.value or '')
            if cell.row == 1:
                continue
            max_len = max(max_len, len(val))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

# ================= SHEET 1: ملخص المشروع =================
ws1 = wb.create_sheet(title="ملخص المشروع والمنصة")
style_sheet(ws1, "Project Master Overview & Architecture")

overview_headers = ["المجال / البند (Category)", "القيمة / التفاصيل (Specification)", "الحالة (Status)", "ملاحظات وتوجيهات أمنية وفنية (Notes)"]
ws1.append([])
ws1.append(overview_headers)
ws1.row_dimensions[3].height = 26

for col_idx in range(1, len(overview_headers) + 1):
    cell = ws1.cell(row=3, column=col_idx)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center")

overview_data = [
    ["اسم البراند والمنصة", "Tourvanto (Excursions & Transfers)", "معتمد ورسمي", "منصة حجز رحلات سياحية وأنشطة وانتقالات بدون فنادق"],
    ["مكان المشروع التخزيني", "D:\\Tourvanto", "محفوظ ومؤكد", "تم الحفظ بالكامل على قرص D لتوفير مساحة C بناء على طلبكم"],
    ["مستودع الكود GitHub", "https://github.com/MohamedElghala/Tourvanto.git", "مربوط ومحدث", "فرع main محدث بآخر طبقات الأمان والبكسلز والتقييمات"],
    ["رابط الموقع الحي Vercel", "https://tourvanto.vercel.app", "منشور ويعمل 200 OK", "جاهز للاستعراض المباشر ومحدث بآخر نسخة"],
    ["رابط لوحة التحكم السرية", "https://tourvanto.vercel.app/#admin", "مخفية ومؤمنة", "تفتح بنافذة المصادقة السرية أو بأيقونة القفل بالفوتر"],
    ["رمز الحماية السري للوحة PIN", "Tourvanto#9900", "مفعل ومحمي", "مطلوب للدخول للـ Dashboard ومتابعة الحجوزات والرحلات"],
    ["اللغات المدعومة", "5 لغات (English, Deutsch, Русский, Français, Italiano)", "مفعل 100%", "قاموس كامل يغطي تفاصيل الرحلات والسياسات والتقييمات"],
    ["العملات المدعومة", "4 عملات (EUR, USD, RUB, GBP)", "مفعل لحظياً", "تحويل مباشر بأسعار الصرف الدقيقة مع حفظ التفضيلات"],
    ["الجمهور المستهدف", "السياح الأجانب في مصر (أوروبا، روسيا، بريطانيا)", "مستهدف", "مصمم لراحة السائح وسهولة الحجز الفوري"],
    ["فحص الأمان Cyber Security", "حماية من XSS، وتشفير Headers، وإخفاء أدوات المطور", "مفحوص ومطبق", "تم تطبيق X-Frame-Options و XSS Sanitization"],
    ["خيار الدفع الحالي النشط", "الدفع عند الوصول كاش (Zero Upfront Cash on Pickup)", "مفعل وجاهز", "أعلى وسيلة تحويل سياحي تضمن ثقة السائح الأجنبي"],
    ["بوابات الدفع الإلكتروني", "مهيأة برمجياً بـ payment-config.js (Stripe & PayPal)", "قيد التجهيز", "في انتظار الفيزا والحساب البنكي لربط المفاتيح"],
    ["قاعدة البيانات والمصادقة", "مهيأة برمجياً بـ db-config.js و auth.js (Gmail, Apple, Outlook)", "جاهزة للربط", "في انتظار عمل الجيميل الرسمي لربط Supabase بدون وهمي"],
    ["أدوات التحليل والبكسلز", "مهيأة بـ analytics.js (GA4, Meta Pixel, TikTok Pixel)", "مربوطة وجاهزة", "تتتبع كل خطوة (ViewContent, Checkout, Purchase)"],
    ["قسم التقييمات التفاعلي", "مفعل مع إمكانية إضافة ريفيو من أي زائر بالنجوم", "مفعل وحي", "يتم حفظ التقييم وعرضه فوراً مع حساب المتوسط"],
    ["البورتفوليو والوسائط", "قسم المعرض والفيديوهات مهيأ وموجود بالصفحة الرئيسية", "جاهز للصور غداً", "يستقبل صور وفيديوهات الجودة العالية فور إحضارها"]
]

for row_idx, row in enumerate(overview_data, start=4):
    ws1.append(row)
    for col_idx in range(1, len(row) + 1):
        c = ws1.cell(row=row_idx, column=col_idx)
        c.font = regular_font
        c.border = cell_border
        if col_idx == 1:
            c.font = bold_font
        if col_idx == 3:
            c.alignment = Alignment(horizontal="center")
            c.fill = highlight_fill

auto_fit_columns(ws1)

# ================= SHEET 2: تحليل المنافسين الشامل =================
ws2 = wb.create_sheet(title="تحليل المنافسين الشامل")
style_sheet(ws2, "Competitor Benchmarking & Strategic Edge")

comp_headers = ["عنصر المقارنة (Feature)", "Viator (Tripadvisor)", "GetYourGuide", "الشركات المحلية التقليدية", "Tourvanto (منصتنا)", "ميزتنا التنافسية وسر التفوق"]
ws2.append([])
ws2.append(comp_headers)
ws2.row_dimensions[3].height = 26

for col_idx in range(1, len(comp_headers) + 1):
    cell = ws2.cell(row=3, column=col_idx)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center")

comp_data = [
    ["عمولة الحجز", "تأخذ من 25% إلى 30% عمولة من المورد", "تأخذ من 20% إلى 30% عمولة باهظة", "غير محددة وتعتمد على السماسرة", "0% عمولة للمنصة (أرباحك مباشرة)", "أسعارنا أرخص للسائح وهوامش ربحك أعلى بنسبة 100%"],
    ["طريقة الدفع للسائح", "دفع بالفيزا إجباري ومقدم 100%", "دفع إلكتروني مسبق فقط", "كاش باليد بدون إيصالات أو ضمان", "خيار مرن: كاش عند الانطلاق أو فيزا/باي بال", "السائح لا يخاف من ضياع أمواله ويثق في الحجز فوراً"],
    ["الدعم والواتساب", "روبوت إيميلات وبطء في الرد", "دعم مراسلة بطيء ويمنع الواتساب", "مكالمات عشوائية بدون نظام", "زر واتساب VIP مباشر مع خدمة عملاء 24/7", "حل مشكلات السائح وتأكيد مواعيده في دقيقة واحدة"],
    ["تنسيق الاستقبال بالفندق", "غامض ويكتفي بإرسال نقطة على الخريطة", "تعليمات نصية قد يخطئ فيها السائح", "تأخير ومشاكل في معرفة الغرف", "حقل مخصص لاسم الفندق ورقم الغرفة وفوتشر دقيق", "السائق والمرشد يعرفان السائح بالاسم ورقم الغرفة بدقة"],
    ["البرامج المخصصة (VIP Builder)", "باقات جاهزة وجامدة لا تقبل التعديل", "باقات ثابتة فقط", "تفاوض يدوي متعب في الشارع", "حاسبة رحلات VIP فورية بالموقع", "يستطيع السائح تفصيل سيارة خاصة وساعات خاصة بضغطة زر"],
    ["اللغات وجودة الترجمة", "ترجمة آلية عامة", "ترجمة آلية للإنجليزية والألمانية", "عربي وإنجليزي ضعيف غالباً", "5 لغات أصلية مدروسة بدقة (EN, DE, RU, FR, IT)", "تخاطب السائح بلغته الأم وخصوصاً السياح الروس والألمان"],
    ["التقييمات والمصداقية", "تقييمات مجمعة قديمة", "تقييمات معقدة الإرسال", "لا توجد أي تقييمات موثقة", "نظام تقييم فوري بالنجوم وعلم الدولة", "أي سائح يشارك تجربته فوراً وتظهر مصداقية المنصة أمام الجميع"],
    ["السرعة وتجربة المستخدم", "موقع ثقيل جداً ومعقد الخطوات", "موقع مزدحم بالفنادق والقطارات", "صفحة فيسبوك بدائية", "منصة فائقة السرعة متخصصة فقط بالرحلات", "تجربة تصفح فاخرة مريحة في الموبايل بدون أي تشتيت"],
    ["تجهيزات التسويق والبكسلز", "تستحوذ على بيانات العميل لنفسها", "تحتكر الجماهير وإعادة الاستهداف", "لا تفهم في البكسل ولا التحليل", "بكسلز متقدمة (تيك توك، فيسبوك، جوجل) ملكك بالكامل", "تستطيع عمل حملات ريتارجيتنج للسياح في شرم والغردقة مباشرة"]
]

for row_idx, row in enumerate(comp_data, start=4):
    ws2.append(row)
    for col_idx in range(1, len(row) + 1):
        c = ws2.cell(row=row_idx, column=col_idx)
        c.font = regular_font
        c.border = cell_border
        if col_idx in [1, 5]:
            c.font = bold_font
        if col_idx == 5:
            c.fill = highlight_fill

auto_fit_columns(ws2)

# ================= SHEET 3: كتالوج الرحلات والأسعار =================
ws3 = wb.create_sheet(title="كتالوج الرحلات والأسعار")
style_sheet(ws3, "Excursions Catalog, Itineraries & Multi-Currency Pricing")

tour_headers = ["كود الرحلة", "اسم الرحلة بالإنجليزية", "التصنيف", "المدة (ساعات)", "سعر اليورو (€)", "سعر الدولار ($)", "سعر الروبل (₽)", "سعر الإسترليني (£)", "التقييم", "عدد التقييمات", "أهم الخدمات المشمولة", "الخدمات المستبعدة", "شارة التسويق المميزة"]
ws3.append([])
ws3.append(tour_headers)
ws3.row_dimensions[3].height = 26

for col_idx in range(1, len(tour_headers) + 1):
    cell = ws3.cell(row=3, column=col_idx)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center")

tours_catalog = [
    ["tour_1", "VIP Mega Quad & Buggy Desert Safari", "Safari", 5, 45, 49, 4433, 38, 4.95, 342, "النقل المكيف من الفندق، بيتش باجي رباعي، ركوب جمال، عشاء بدوي بوفيه، عروض فلكلور", "الشال البدوي والنظارات الواقية (تباع رمزياً)", "🔥 Likely to sell out soon"],
    ["tour_2", "Orange Bay Island & Dolphin Luxury Cruise", "Sea", 7, 55, 59, 5418, 47, 4.98, 489, "يخت فاخر، جزيرة أورانج باي ساعتين، مشاهدة الدلافين، معدات سنوركلينج، غداء سي فود", "الصور الفوتوغرافية الخاصة بالمصور باليخت", "🏆 Badge of Excellence 2026"],
    ["tour_3", "Private Cairo & Pyramids Day Expedition", "Culture", 14, 120, 130, 11820, 102, 4.92, 215, "سيارة خاصة مرسيدس بسائق، مرشد مصريات خاص، تذاكر الأهرامات وأبو الهول، غداء فاخر", "دخول داخل الهرم الأكبر (اختياري إضافي)", "👑 Top Rated Cultural Tour"],
    ["tour_4", "Luxor & Valley of the Kings Majestic Discovery", "Culture", 15, 110, 119, 10835, 94, 4.96, 178, "انتقال خاص ومريح، معبد الكرنك، وادي الملوك 3 مقابر، معبد حتشبسوت، غداء نيلي", "مقبرة توت عنخ آمون (تذكرة خاصة)", "✨ Best Seller"],
    ["tour_5", "Red Sea Intro Scuba Diving Experience", "Sea", 6, 65, 70, 6403, 55, 4.97, 156, "غطستين مع مدربين معتمدين PADI، معدات الغطس كاملة، غداء ومشروبات طوال اليوم باليخت", "شهادة غطس رسمية (تحتاج دورة كاملة)", "🤿 100% Beginner Friendly"],
    ["tour_6", "Private Luxury VIP Airport Transfer", "Transfers", 1, 35, 38, 3448, 30, 4.99, 520, "سيارة VIP فاخرة خاصة، استقبال بالاسم في المطار، مساعدة في الحقائب، مياه معدنية وواي فاي", "المشروبات الكحولية", "⚡ Instant Pick-up Confirmation"]
]

for row_idx, row in enumerate(tours_catalog, start=4):
    ws3.append(row)
    for col_idx in range(1, len(row) + 1):
        c = ws3.cell(row=row_idx, column=col_idx)
        c.font = regular_font
        c.border = cell_border
        if col_idx in [1, 2]:
            c.font = bold_font
        if col_idx in [5, 6, 7, 8]:
            c.alignment = Alignment(horizontal="right")
        if col_idx in [4, 9, 10]:
            c.alignment = Alignment(horizontal="center")

auto_fit_columns(ws3)

# ================= SHEET 4: دليل البكسلز والتسويق =================
ws4 = wb.create_sheet(title="دليل البكسلز والتسويق")
style_sheet(ws4, "Marketing Pixels, Tracking Architecture & Audience Funnels")

pixel_headers = ["المنصة الإعلانية (Platform)", "اسم البكسل / المعرف المطلوب", "مكان وضعه في الكود", "الحدث المتتبع (Event Name)", "متى يعمل الحدث؟ (Trigger)", "البيانات المرسلة (Data Payload)", "الفائدة الإعلانية والجمهور المجمع"]
ws4.append([])
ws4.append(pixel_headers)
ws4.row_dimensions[3].height = 26

for col_idx in range(1, len(pixel_headers) + 1):
    cell = ws4.cell(row=3, column=col_idx)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center")

pixel_data = [
    ["Google Analytics 4 (GA4)", "G-XXXXXXXXXX", "D:\\Tourvanto\\analytics.js", "page_view", "زيارة أي صفحة في الموقع", "URL, Page Title, User Agent", "معرفة الدول ومصادر الزيارات وسلوك الزوار اليومي"],
    ["Google Analytics 4 (GA4)", "G-XXXXXXXXXX", "D:\\Tourvanto\\analytics.js", "view_item", "عند فتح تفاصيل أي رحلة", "Tour ID, Tour Name, Category, Price EUR", "بناء جمهور لمهتمي رحلات السفاري أو الغوص وإعادة استهدافهم"],
    ["Google Analytics 4 (GA4)", "G-XXXXXXXXXX", "D:\\Tourvanto\\analytics.js", "begin_checkout", "عند الضغط على احجز الآن وفتح الاستمارة", "Tour Title, Total Amount, Currency", "تتبع نسبة التسرب في قمع الحجز ومعالجة المشاكل"],
    ["Google Analytics 4 (GA4)", "G-XXXXXXXXXX", "D:\\Tourvanto\\analytics.js", "purchase", "عند تأكيد الحجز وإصدار الفاوتشر", "Transaction ID, Total Value, Payment Method", "قياس عائد الاستثمار ROAS للحملات الإعلانية المدفوعة"],
    ["Meta Pixel (فيسبوك وإنستغرام)", "15 أو 16 رقم (Pixel ID)", "D:\\Tourvanto\\analytics.js", "PageView", "فتح أي صفحة في الموقع", "URL, Referrer", "تجميع كل زوار الموقع في Custom Audience بحساب الإعلانات"],
    ["Meta Pixel (فيسبوك وإنستغرام)", "15 أو 16 رقم (Pixel ID)", "D:\\Tourvanto\\analytics.js", "ViewContent", "عند تصفح رحلة معينة", "content_name, content_category, value, currency", "عمل Dynamic Product Ads لكل سائح بالرحلة التي شاهدها"],
    ["Meta Pixel (فيسبوك وإنستغرام)", "15 أو 16 رقم (Pixel ID)", "D:\\Tourvanto\\analytics.js", "InitiateCheckout", "عند بدء تعبئة استمارة الحجز", "content_name, value, currency", "إعادة استهداف السائحين الذين بدأوا الحجز ولم يكملوه"],
    ["Meta Pixel (فيسبوك وإنستغرام)", "15 أو 16 رقم (Pixel ID)", "D:\\Tourvanto\\analytics.js", "Purchase / Lead", "عند الضغط على تأكيد الحجز بنجاح", "transaction_id, value, payment_method", "تحسين خوارزمية فيسبوك لتبحث عن عملاء مشابهين (Lookalike 1%)"],
    ["TikTok Pixel", "كود تعريف التيك توك", "D:\\Tourvanto\\analytics.js", "Pageview", "دخول السائح للموقع من التيك توك", "Page URL, Timestamp", "بناء جمهور التيك توك وتتبع سرعة استجابة السائحين بالفيديوهات"],
    ["TikTok Pixel", "كود تعريف التيك توك", "D:\\Tourvanto\\analytics.js", "ViewContent", "مشاهدة رحلة الغطس أو السفاري", "content_id, content_name, value", "إظهار فيديوهات إضافية للرحلة للسائح داخل التيك توك"],
    ["TikTok Pixel", "كود تعريف التيك توك", "D:\\Tourvanto\\analytics.js", "PlaceAnOrder", "إتمام حجز السائح", "order_id, value, currency", "تحسين حملات التيك توك للوصول إلى السياح المتواجدين في مصر حالياً"]
]

for row_idx, row in enumerate(pixel_data, start=4):
    ws4.append(row)
    for col_idx in range(1, len(row) + 1):
        c = ws4.cell(row=row_idx, column=col_idx)
        c.font = regular_font
        c.border = cell_border
        if col_idx in [1, 4]:
            c.font = bold_font

auto_fit_columns(ws4)

# ================= SHEET 5: خريطة الخطوات والربط الفني =================
ws5 = wb.create_sheet(title="خريطة الخطوات والربط الفني")
style_sheet(ws5, "Master Launch Roadmap & Step-by-Step Execution Plan")

plan_headers = ["رقم الخطوة", "المهمة والهدف (Action Item)", "المسؤولية", "الحالة الحالية", "الخطوات التنفيذية بالتفصيل (How to execute)"]
ws5.append([])
ws5.append(plan_headers)
ws5.row_dimensions[3].height = 26

for col_idx in range(1, len(plan_headers) + 1):
    cell = ws5.cell(row=3, column=col_idx)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center")

roadmap_data = [
    ["1", "إنشاء إيميل جيميل رسمي مخصص للمشروع (مثل tourvanto.official@gmail.com)", "العميل", "خطوة أولى مطلوبة", "إنشاء بريد Gmail جديد يكون المركز الموحد لربط قاعدة البيانات، الحسابات البنكية، والبكسلز."],
    ["2", "ربط قاعدة البيانات الرسمية (Supabase) بالإيميل الجديد بدون قواعد وهمية", "الوكيل والعميل", "مهيأة برمجياً بـ db-config.js", "فتح حساب مجاني في supabase.com بالإيميل الرسمي، والحصول على Project URL و Anon Key ووضعهما في db-config.js."],
    ["3", "حجز الدومين الرسمي tourvanto.com وربطه بـ Vercel DNS", "العميل والوكيل", "في الانتظار", "شراء الدومين من Namecheap أو Cloudflare وربطه بلوحة Vercel Domains ليصبح الموقع .com رسمي."],
    ["4", "تحديث أرقام الواتساب وخدمة العملاء الحقيقية", "العميل", "مهيأ برقم تجريبي", "استبدال رقم +20 100 000 0000 بالرقم الفعلي للشركة لاستقبال حجوزات السائحين الفورية."],
    ["5", "استخراج ووضع معرفات البكسلز (GA4, Meta, TikTok) بحساب الإيميل الجديد", "العميل والوكيل", "الكود مهيأ بـ analytics.js", "فتح Meta Events Manager و TikTok Ads Manager بنفس الإيميل، ونسخ المعرفات ولصقها في analytics.js."],
    ["6", "ربط بوابات الدفع بالفيزا وباي بال (Stripe & PayPal)", "العميل والوكيل", "مهيأ بـ payment-config.js", "بمجرد استخراج فيزا بنكية تجارية، تفعيل حساب Stripe ووضع Publishable Key ليعمل الدفع المباشر فوراً."],
    ["7", "إحضار صور وفيديوهات البورتفوليو غداً ورفعها بالموقع", "العميل", "المعرض مهيأ وجاهز لاستقبالها", "سنقوم برفع الصور عالية الدقة ومقاطع الفيديو 4K في قسم البورتفوليو وتحديث كروت الرحلات بها."]
]

for row_idx, row in enumerate(roadmap_data, start=4):
    ws5.append(row)
    for col_idx in range(1, len(row) + 1):
        c = ws5.cell(row=row_idx, column=col_idx)
        c.font = regular_font
        c.border = cell_border
        if col_idx in [1, 2]:
            c.font = bold_font
        if col_idx == 4:
            c.fill = highlight_fill

auto_fit_columns(ws5)

# Save workbook
output_path = r"D:\Tourvanto\Tourvanto_Project_Full_Data.xlsx"
wb.save(output_path)
print(f"Workbook successfully saved to: {output_path}")
